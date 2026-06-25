from __future__ import annotations

import argparse
import dataclasses
import json
import re
import shutil
from pathlib import Path
from typing import Iterable, List, Optional, Sequence, Tuple

import fitz
from docx import Document


DEFAULT_OUTPUT_ROOT = Path(r"D:\Codex_output")
ILLEGAL_FILENAME_CHARS = r'[<>:"/\\|?*]'
HEADING_NUMBER_RE = re.compile(r"^(\d+\.\d+(?:\.\d+){0,4})[\s\u3000]+(.+)$")
CHINESE_HEADING_RE = re.compile(r"^第[一二三四五六七八九十百千0-9]+[编章节部分篇](?:[\s\u3000].*|$)")
TOC_LEADER_RE = re.compile(r"\.{5,}\s*\d+\s*$")
CHAPTER_SUMMARY_RE = re.compile(r"^第[一二三四五六七八九十百千0-9]+[编章节部分篇](主要|重点|概述|介绍|从)")
PAREN_ENUM_RE = re.compile(r"^[（(]\s*\d+\s*[)）]")
GUIDE_TITLE_RE = re.compile(r".{2,}指南$")
TABLE_TITLE_RE = re.compile(r"^表\s*\d+")
TABLE_SOURCE_RE = re.compile(r"^(数据来源|来源)[:：]")
DATE_LED_SENTENCE_RE = re.compile(r"^\d{4}\s*年(?:\s*\d{1,2}\s*月(?:\s*\d{1,2}\s*日)?)?(起|，|,|度|内|末|初)")
SENTENCE_PUNCT_RE = re.compile(r"[，；。！？]")
BRACKET_LEAD_RE = re.compile(r"^[\[【(（]")
FORMULA_LEAD_RE = re.compile(r"^\d+\s*[+\-*/]")
CIRCLED_ENUM_RE = re.compile(r"^[①②③④⑤⑥⑦⑧⑨⑩⑪⑫⑬⑭⑮⑯⑰⑱⑲⑳]")
LATIN_ENUM_RE = re.compile(r"^[A-Z][.．、]")
APPENDIX_ITEM_RE = re.compile(r"^[一二三四五六七八九十]+[、.．]\s*.+$")
APPENDIX_LETTER_HEADING_RE = re.compile(r"^附录\s*[A-ZＡ-Ｚ](?:[\s\u3000]+|[：:])?.+$")
APPENDIX_CHINESE_HEADING_RE = re.compile(r"^附录[一二三四五六七八九十]+(?:[\s\u3000]+|[：:])?.+$")
PERSON_NAME_LINE_RE = re.compile(r"^[\u4e00-\u9fff]{2,4}(?:[\s\u3000]+[\u4e00-\u9fff]{2,4}){1,}$")
FORMAL_SHORT_HEADINGS = {"前言", "目录", "附录"}
HEADING_SUFFIX_FRAGMENTS = {"业", "表", "图", "览表", "税率表"}


def is_appendix_heading(text: str) -> bool:
    normalized = re.sub(r"\s+", " ", text.strip())
    return bool(
        APPENDIX_LETTER_HEADING_RE.match(normalized)
        or APPENDIX_CHINESE_HEADING_RE.match(normalized)
    )


def heading_remainder_looks_like_body(text: str) -> bool:
    normalized = re.sub(r"\s+", " ", text.strip())
    if not normalized:
        return False

    if SENTENCE_PUNCT_RE.search(normalized):
        return True

    if FORMULA_LEAD_RE.match(normalized):
        return True

    if normalized.startswith(("+", "-", "*", "/")):
        return True

    if BRACKET_LEAD_RE.match(normalized):
        return True

    if re.match(r"^\d{4}\s*年", normalized):
        return True

    if re.match(r"^\d{1,2}\s*月", normalized):
        return True

    if re.match(r"^\d+\s*(个|款|项|条|年|月|日)", normalized):
        return True

    if re.match(r"^\d+(?:\.\d+)?\s*(亿|万|元|美元|欧元|新台币|%)", normalized):
        return True

    return False


def is_formal_short_heading(text: str) -> bool:
    normalized = re.sub(r"\s+", "", text.strip())
    return normalized in FORMAL_SHORT_HEADINGS


def looks_like_person_name_line(text: str) -> bool:
    normalized = re.sub(r"\s+", " ", text.strip())
    if not PERSON_NAME_LINE_RE.fullmatch(normalized):
        return False

    parts = [part for part in re.split(r"[\s\u3000]+", normalized) if part]
    return len(parts) >= 2


def looks_like_compact_person_names(text: str) -> bool:
    compact = re.sub(r"\s+", "", text.strip())
    return bool(re.fullmatch(r"[\u4e00-\u9fff]{4,8}", compact))


@dataclasses.dataclass
class Block:
    kind: str
    text: str
    level: Optional[int] = None
    bbox: Optional[Tuple[float, float, float, float]] = None


@dataclasses.dataclass
class Section:
    title: str
    heading_path: List[Block]
    content: List[Block]


@dataclasses.dataclass
class TableNarrative:
    bbox: Tuple[float, float, float, float]
    text: str
    consumed_bboxes: List[Tuple[float, float, float, float]]


def sanitize_filename(title: str) -> str:
    cleaned = title.strip().replace("/", "_").replace("\\", "_").replace(":", "_")
    cleaned = re.sub(r'[<>\"|?*]', " ", cleaned)
    cleaned = cleaned.replace("\n", " ").replace("\r", " ")
    cleaned = cleaned.replace('"', "")
    cleaned = re.sub(r"\s+", " ", cleaned)
    cleaned = re.sub(r" ?_+ ?", "_", cleaned)
    cleaned = cleaned.strip(" ._")
    return cleaned or "untitled-section"


def build_measurement_clause(header: str, value: str) -> str:
    compact_header = re.sub(r"\s+", "", header)
    if "兑换" in compact_header:
        left, right = compact_header.split("兑换", 1)
        left = re.sub(r"(\d)([\u4e00-\u9fff])", r"\1 \2", left)
        right = re.sub(r"([\u4e00-\u9fff])(\d)", r"\1 \2", right)
        if right:
            return f"{left}兑换 {value} {right}"
        return f"{left}兑换 {value}"
    return f"{header}{value}"


def narrate_table(
    table: Sequence[Sequence[object]],
    title: Optional[str] = None,
    source: Optional[str] = None,
) -> str:
    rows = [[str(cell).strip() for cell in row] for row in table if any(str(cell).strip() for cell in row)]
    if not rows:
        return "表格为空，未提取到可叙述的数据。"

    width = max(len(row) for row in rows)
    normalized = [row + [""] * (width - len(row)) for row in rows]
    headers = normalized[0]
    header_names = [header or f"第{i + 1}列" for i, header in enumerate(headers)]

    if title:
        parts = [title]
    else:
        parts = [f"表格包含 {width} 列、{max(len(normalized) - 1, 0)} 行数据。"]

    year_like_first_column = headers and (headers[0] in {"年份", "年度", "时间"} or all(re.fullmatch(r"\d{4}", row[0]) for row in normalized[1:] if row[0]))
    measurable_columns = len(header_names) >= 2 and all(any(char.isdigit() for char in name) or "兑换" in name or "汇率" in name for name in header_names[1:])

    if year_like_first_column and measurable_columns:
        for row in normalized[1:]:
            label = row[0] or "该年度"
            clauses = []
            for header, value in zip(header_names[1:], row[1:]):
                if not value:
                    continue
                clauses.append(build_measurement_clause(header, value))
            sentence = "；".join(clauses) if clauses else "无可提取数据"
            parts.append(f"{label} 年 {sentence}。")
    else:
        parts.append(f"列标题依次为：{'、'.join(header_names)}。")
        for row_index, row in enumerate(normalized[1:], start=1):
            assignments = []
            for header, value in zip(header_names, row):
                assignments.append(f"{header}为“{value or '空白'}”")
            parts.append(f"第 {row_index} 行中，" + "，".join(assignments) + "。")

    if source:
        parts.append(source)

    return "\n".join(parts)


def split_sections(blocks: Sequence[Block]) -> List[Section]:
    sections: List[Section] = []
    heading_stack: List[Block] = []
    current_section: Optional[Section] = None

    for block in blocks:
        if block.kind == "heading":
            level = block.level or 1
            heading_stack = [item for item in heading_stack if (item.level or 1) < level]
            heading_stack.append(block)
            current_section = Section(
                title=block.text,
                heading_path=list(heading_stack),
                content=[],
            )
            sections.append(current_section)
            continue

        if current_section is not None and block.text.strip():
            current_section.content.append(block)

    return sections


def is_numbered_heading(text: str) -> bool:
    normalized = re.sub(r"\s+", " ", text.strip())
    return bool(HEADING_NUMBER_RE.match(normalized) or CHINESE_HEADING_RE.match(normalized))


def should_merge_heading_pair(previous: Block, current: Block) -> bool:
    if previous.kind != "heading" or current.kind != "heading":
        return False

    previous_text = re.sub(r"\s+", " ", previous.text.strip())
    current_text = re.sub(r"\s+", " ", current.text.strip())
    if not previous_text or not current_text:
        return False

    if is_formal_short_heading(previous_text + current_text):
        return True

    if is_appendix_heading(previous_text):
        if is_numbered_heading(current_text):
            return False
        if PAREN_ENUM_RE.match(current_text) or CIRCLED_ENUM_RE.match(current_text):
            return False
        if heading_remainder_looks_like_body(current_text):
            return False
        return len(re.sub(r"\s+", "", current_text)) <= 3

    if not CHINESE_HEADING_RE.match(previous_text):
        return False
    if is_numbered_heading(current_text):
        return False
    if PAREN_ENUM_RE.match(current_text):
        return False
    if CIRCLED_ENUM_RE.match(current_text):
        return False
    if is_guide_title(current_text):
        return False
    if heading_remainder_looks_like_body(current_text):
        return False
    if len(current_text) > 20:
        return False

    return True


def merge_multiline_headings(blocks: Sequence[Block]) -> List[Block]:
    merged: List[Block] = []
    for block in blocks:
        if merged and should_merge_heading_pair(merged[-1], block):
            previous = merged[-1]
            previous_text = previous.text.strip()
            current_text = block.text.strip()
            merged_text = (
                previous_text + current_text
                if is_formal_short_heading(previous_text + current_text)
                else f"{previous_text} {current_text}"
            )
            merged[-1] = Block(
                "heading",
                merged_text,
                level=previous.level,
                bbox=previous.bbox,
            )
            continue
        merged.append(block)
    return merged


def merge_heading_fragments(blocks: Sequence[Block]) -> List[Block]:
    merged: List[Block] = []
    for block in blocks:
        if merged:
            previous = merged[-1]
            previous_text = previous.text.strip()
            current_text = block.text.strip()
            compact_current = re.sub(r"\s+", "", current_text)
            if (
                previous.kind == "heading"
                and is_appendix_heading(previous_text)
                and block.kind == "paragraph"
                and compact_current in HEADING_SUFFIX_FRAGMENTS
                and compact_current not in FORMAL_SHORT_HEADINGS
                and not heading_remainder_looks_like_body(current_text)
            ):
                merged[-1] = Block(
                    "heading",
                    f"{previous_text}{compact_current}",
                    level=previous.level,
                    bbox=previous.bbox,
                )
                continue
        merged.append(block)
    return merged


def drop_name_only_headings(blocks: Sequence[Block]) -> List[Block]:
    filtered: List[Block] = []
    parent_headings: List[Block] = []

    for block in blocks:
        if block.kind == "heading":
            level = block.level or 1
            parent_headings = [item for item in parent_headings if (item.level or 1) < level]
            parent_title = parent_headings[-1].text if parent_headings else ""
            if (looks_like_person_name_line(block.text) or looks_like_compact_person_names(block.text)) and parent_title in {"参考文献"}:
                filtered.append(Block("paragraph", re.sub(r"\s+", "", block.text.strip()), bbox=block.bbox))
                continue
            parent_headings.append(block)
        filtered.append(block)

    return filtered


def promote_appendix_items(blocks: Sequence[Block]) -> List[Block]:
    promoted: List[Block] = []
    in_appendix = False

    for block in blocks:
        normalized = re.sub(r"\s+", "", block.text.strip())
        if block.kind == "heading" and normalized == "附录":
            in_appendix = True
            promoted.append(block)
            continue

        if in_appendix and APPENDIX_ITEM_RE.match(block.text.strip()):
            promoted.append(Block("heading", block.text, level=2, bbox=block.bbox))
            continue

        if in_appendix and block.kind == "heading" and (block.level or 1) <= 1 and normalized != "附录":
            in_appendix = False

        promoted.append(block)

    return promoted


def is_guide_title(text: str) -> bool:
    normalized = re.sub(r"\s+", "", text.strip())
    return bool(normalized and GUIDE_TITLE_RE.fullmatch(normalized))


def find_guide_title(blocks: Sequence[Block]) -> Optional[str]:
    normalized_blocks = [re.sub(r"\s+", "", block.text.strip()) for block in blocks if block.text.strip()]
    for index, text in enumerate(normalized_blocks[:8]):
        if not text.endswith("指南"):
            continue
        combined = text
        back_index = index - 1
        while back_index >= 0:
            previous = normalized_blocks[back_index]
            if not previous or len(previous) > 30 or previous.endswith(("组", "司", "局")):
                break
            combined = previous + combined
            if combined.endswith("指南"):
                back_index -= 1
                continue
            break
        if GUIDE_TITLE_RE.fullmatch(combined):
            return combined
        if is_guide_title(text):
            return text
    return None


def filter_redundant_guide_lines(blocks: Sequence[Block]) -> List[Block]:
    guide_title = find_guide_title(blocks)
    if not guide_title:
        return list(blocks)

    kept: List[Block] = []
    seen_first = False
    for block in blocks:
        normalized = re.sub(r"\s+", "", block.text.strip())
        if normalized == guide_title:
            if not seen_first:
                kept.append(Block(block.kind, guide_title, level=block.level, bbox=block.bbox))
                seen_first = True
            continue
        kept.append(block)
    return kept


def has_tesseract() -> bool:
    return shutil.which("tesseract") is not None


def overlap_ratio(bbox_a: Tuple[float, float, float, float], bbox_b: Tuple[float, float, float, float]) -> float:
    ax0, ay0, ax1, ay1 = bbox_a
    bx0, by0, bx1, by1 = bbox_b
    ix0 = max(ax0, bx0)
    iy0 = max(ay0, by0)
    ix1 = min(ax1, bx1)
    iy1 = min(ay1, by1)
    if ix1 <= ix0 or iy1 <= iy0:
        return 0.0
    intersection = (ix1 - ix0) * (iy1 - iy0)
    area = max((ax1 - ax0) * (ay1 - ay0), 1.0)
    return intersection / area


def horizontal_overlap(bbox_a: Tuple[float, float, float, float], bbox_b: Tuple[float, float, float, float]) -> float:
    ax0, _, ax1, _ = bbox_a
    bx0, _, bx1, _ = bbox_b
    intersection = min(ax1, bx1) - max(ax0, bx0)
    if intersection <= 0:
        return 0.0
    width = max(min(ax1 - ax0, bx1 - bx0), 1.0)
    return intersection / width


def looks_like_heading(text: str, font_size: Optional[float], body_font_size: float) -> Optional[int]:
    normalized = re.sub(r"\s+", " ", text.strip())
    if not normalized:
        return None

    if TOC_LEADER_RE.search(normalized):
        return None

    if is_formal_short_heading(normalized):
        return 1

    if is_appendix_heading(normalized):
        return 1

    if PAREN_ENUM_RE.match(normalized):
        return None

    if CIRCLED_ENUM_RE.match(normalized):
        return None

    if LATIN_ENUM_RE.match(normalized):
        return None

    if is_guide_title(normalized):
        return None

    if looks_like_person_name_line(normalized):
        return None

    if DATE_LED_SENTENCE_RE.match(normalized):
        return None

    if CHAPTER_SUMMARY_RE.match(normalized):
        return None

    match = HEADING_NUMBER_RE.match(normalized)
    if match:
        if heading_remainder_looks_like_body(match.group(2)):
            return None
        return min(match.group(1).count(".") + 1, 6)

    if CHINESE_HEADING_RE.match(normalized):
        return 1

    if normalized.endswith(("。", "；", "：", "，", ".", ";", ":", ",")):
        return None

    if len(normalized) > 40:
        return None

    if heading_remainder_looks_like_body(normalized):
        return None

    compact = re.sub(r"\s+", "", normalized)
    if compact not in FORMAL_SHORT_HEADINGS and len(compact) <= 2:
        return None

    if font_size is None:
        if len(normalized) > 12 or re.search(r"\d", normalized):
            return None
        return 3

    if len(normalized) > 12 or re.search(r"\d", normalized):
        return None

    if font_size >= body_font_size + 5:
        return 1
    if font_size >= body_font_size + 3:
        return 2
    if font_size >= body_font_size + 1.5:
        return 3

    return None


def line_blocks_from_text(text: str) -> List[Block]:
    blocks: List[Block] = []
    for raw_line in text.splitlines():
        line = raw_line.strip()
        if not line:
            continue
        level = looks_like_heading(line, None, 12.0)
        kind = "heading" if level else "paragraph"
        blocks.append(Block(kind, line, level=level))
    return blocks


def extract_text_and_dict(
    page: fitz.Page,
    ocr_lang: str,
    min_direct_chars: int,
    force_ocr: bool = False,
) -> Tuple[str, dict, bool]:
    direct_text = page.get_text("text").strip()
    direct_chars = len(re.sub(r"\s+", "", direct_text))
    use_ocr = force_ocr or direct_chars < min_direct_chars

    if use_ocr:
        if not has_tesseract():
            raise RuntimeError(
                "检测到该页可能需要 OCR，但系统中未找到 tesseract。"
                "请先安装 Tesseract OCR 并确保命令 `tesseract` 可用。"
            )
        textpage = page.get_textpage_ocr(language=ocr_lang, dpi=300, full=True)
        text = page.get_text("text", textpage=textpage).strip()
        page_dict = page.get_text("dict", textpage=textpage)
        return text, page_dict, True

    return direct_text, page.get_text("dict"), False


def page_body_font_size(page_dict: dict) -> float:
    sizes: List[float] = []
    for block in page_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        for line in block.get("lines", []):
            for span in line.get("spans", []):
                text = span.get("text", "").strip()
                if text:
                    sizes.append(round(float(span.get("size", 12.0)), 1))

    if not sizes:
        return 12.0

    counts = {}
    for size in sizes:
        counts[size] = counts.get(size, 0) + 1
    return max(counts, key=counts.get)


def collect_text_blocks(page_dict: dict) -> List[Tuple[Tuple[float, float, float, float], str]]:
    text_blocks: List[Tuple[Tuple[float, float, float, float], str]] = []
    for block in page_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        bbox = tuple(block.get("bbox", (0.0, 0.0, 0.0, 0.0)))
        lines = []
        for line in block.get("lines", []):
            parts = []
            for span in line.get("spans", []):
                if span.get("text", "").strip():
                    parts.append(span.get("text", ""))
            line_text = "".join(parts).strip()
            if line_text:
                lines.append(line_text)
        text = "\n".join(lines).strip()
        if text:
            text_blocks.append((bbox, text))
    return text_blocks


def extract_tables(page: fitz.Page, page_dict: dict) -> List[TableNarrative]:
    tables: List[TableNarrative] = []
    text_blocks = collect_text_blocks(page_dict)
    finder = page.find_tables()
    for table in finder.tables:
        bbox = tuple(table.bbox)
        title = None
        source = None
        consumed_bboxes: List[Tuple[float, float, float, float]] = []

        for text_bbox, text in text_blocks:
            vertical_gap = bbox[1] - text_bbox[3]
            if 0 <= vertical_gap <= 40 and horizontal_overlap(text_bbox, bbox) >= 0.5:
                if TABLE_TITLE_RE.match(text) and title is None:
                    title = text
                    consumed_bboxes.append(text_bbox)
                elif TABLE_SOURCE_RE.match(text) and source is None:
                    source = text
                    consumed_bboxes.append(text_bbox)
            if 0 <= text_bbox[1] - bbox[3] <= 30 and horizontal_overlap(text_bbox, bbox) >= 0.5:
                if TABLE_SOURCE_RE.match(text) and source is None:
                    source = text
                    consumed_bboxes.append(text_bbox)

        tables.append(
            TableNarrative(
                bbox=bbox,
                text=narrate_table(table.extract(), title=title, source=source),
                consumed_bboxes=consumed_bboxes,
            )
        )
    return tables


def extract_page_blocks(page: fitz.Page, ocr_lang: str, min_direct_chars: int) -> Tuple[List[Block], bool]:
    page_text, page_dict, used_ocr = extract_text_and_dict(page, ocr_lang, min_direct_chars)
    body_font = page_body_font_size(page_dict)
    table_entries = extract_tables(page, page_dict)

    extracted: List[Tuple[float, float, int, Block]] = []
    consumed_table_bboxes = []
    for table_index, entry in enumerate(table_entries):
        extracted.append((entry.bbox[1], entry.bbox[0], table_index, Block("table", entry.text, bbox=entry.bbox)))
        consumed_table_bboxes.extend(entry.consumed_bboxes)

    order_index = len(extracted)
    for block in page_dict.get("blocks", []):
        if block.get("type") != 0:
            continue
        bbox = tuple(block.get("bbox", (0.0, 0.0, 0.0, 0.0)))
        if any(overlap_ratio(bbox, entry.bbox) >= 0.3 for entry in table_entries):
            continue
        if any(overlap_ratio(bbox, consumed_bbox) >= 0.8 for consumed_bbox in consumed_table_bboxes):
            continue

        texts: List[str] = []
        font_sizes: List[float] = []
        for line in block.get("lines", []):
            line_parts = []
            for span in line.get("spans", []):
                span_text = span.get("text", "")
                if span_text.strip():
                    line_parts.append(span_text)
                    font_sizes.append(float(span.get("size", body_font)))
            text_line = "".join(line_parts).strip()
            if text_line:
                texts.append(text_line)

        text = "\n".join(texts).strip()
        if not text:
            continue

        font_size = max(font_sizes) if font_sizes else body_font
        level = looks_like_heading(text, font_size, body_font)
        kind = "heading" if level else "paragraph"
        extracted.append((bbox[1], bbox[0], order_index, Block(kind, text, level=level, bbox=bbox)))
        order_index += 1

    if not extracted and page_text:
        return line_blocks_from_text(page_text), used_ocr

    extracted.sort(key=lambda item: (round(item[0], 1), round(item[1], 1), item[2]))
    return [item[3] for item in extracted], used_ocr


def make_unique_path(output_dir: Path, title: str, used_names: dict) -> Path:
    stem = sanitize_filename(title)
    counter = used_names.get(stem, 0) + 1
    used_names[stem] = counter
    if counter == 1:
        name = stem
    else:
        name = f"{stem}-{counter}"
    return output_dir / f"{name}.docx"


def write_section_docx(section: Section, path: Path, guide_title: Optional[str] = None) -> None:
    document = Document()
    if guide_title:
        document.add_paragraph(guide_title)
    for heading in section.heading_path:
        level = min(max(heading.level or 1, 1), 9)
        document.add_heading(heading.text, level=level)

    for block in section.content:
        if block.kind == "paragraph":
            document.add_paragraph(block.text)
        elif block.kind == "table":
            document.add_paragraph(block.text)

    document.save(path)


def process_pdf(input_pdf: Path, output_dir: Path, ocr_lang: str, min_direct_chars: int) -> List[dict]:
    output_dir.mkdir(parents=True, exist_ok=True)
    doc = fitz.open(input_pdf)
    all_blocks: List[Block] = []
    used_ocr_pages: List[int] = []

    for page_index, page in enumerate(doc, start=1):
        blocks, used_ocr = extract_page_blocks(page, ocr_lang, min_direct_chars)
        all_blocks.extend(blocks)
        if used_ocr:
            used_ocr_pages.append(page_index)

    all_blocks = filter_redundant_guide_lines(all_blocks)
    all_blocks = merge_multiline_headings(all_blocks)
    all_blocks = merge_heading_fragments(all_blocks)
    all_blocks = promote_appendix_items(all_blocks)
    all_blocks = drop_name_only_headings(all_blocks)
    guide_title = find_guide_title(all_blocks)
    sections = split_sections(all_blocks)
    if not sections:
        raise RuntimeError("未识别到任何可导出的标题分节，无法生成按节拆分的 Word 文件。")

    manifest = []
    used_names = {}
    for section in sections:
        output_path = make_unique_path(output_dir, section.title, used_names)
        write_section_docx(section, output_path, guide_title=guide_title)
        manifest.append(
            {
                "title": section.title,
                "heading_path": [heading.text for heading in section.heading_path],
                "output_path": str(output_path),
                "content_blocks": len(section.content),
            }
        )

    manifest_path = output_dir / "manifest.json"
    manifest_path.write_text(
        json.dumps(
            {
                "input_pdf": str(input_pdf),
                "ocr_language": ocr_lang,
                "used_ocr_pages": used_ocr_pages,
                "sections": manifest,
            },
            ensure_ascii=False,
            indent=2,
        ),
        encoding="utf-8",
    )
    return manifest


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(
        description="OCR a PDF, narrate tables as prose, and export one DOCX per detected section."
    )
    parser.add_argument("input_pdf", type=Path, help="Path to the source PDF.")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=None,
        help="Directory for generated DOCX files. Defaults to D:\\Codex_output\\<pdf-stem>-sections.",
    )
    parser.add_argument(
        "--ocr-lang",
        default="chi_sim+eng",
        help="Tesseract language pack string used for OCR pages.",
    )
    parser.add_argument(
        "--min-direct-chars",
        type=int,
        default=80,
        help="Pages with fewer extracted characters than this threshold trigger OCR.",
    )
    return parser


def main(argv: Optional[Sequence[str]] = None) -> int:
    parser = build_parser()
    args = parser.parse_args(argv)
    input_pdf = args.input_pdf.resolve()
    output_dir = args.output_dir or DEFAULT_OUTPUT_ROOT / f"{input_pdf.stem}-sections"

    process_pdf(input_pdf, output_dir.resolve(), args.ocr_lang, args.min_direct_chars)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
